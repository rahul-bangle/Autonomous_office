import os
from bs4 import BeautifulSoup
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def convert_html_to_xlsx(html_path, xlsx_path, csv_path):
    print(f"Reading HTML from: {html_path}")
    with open(html_path, 'r', encoding='utf-8') as f:
        soup = BeautifulSoup(f, 'html.parser')

    title = soup.find('h1').text.strip() if soup.find('h1') else "Origin Command Execution Plan"
    meta = soup.find('div', class_='meta').text.strip() if soup.find('div', class_='meta') else ""

    table = soup.find('table')
    if not table:
        print("Error: No table found in HTML.")
        return

    # Extract headers
    headers = [th.text.strip() for th in table.find('thead').find_all('th')]
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Execution Plan"

    # Style definitions
    header_fill = PatternFill(start_color="1A1A1A", end_color="1A1A1A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    phase_fill = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")
    phase_font = Font(bold=True)
    done_font = Font(color="1A7A1A")
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    # Add Title and Meta
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(row=1, column=1).font = Font(size=14, bold=True)
    
    ws.append([meta])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws.cell(row=2, column=1).font = Font(italic=True, size=10)
    
    ws.append([]) # Spacer

    # Add Table Headers
    header_row_idx = 4
    ws.append(headers)
    for col_idx, cell in enumerate(ws[header_row_idx], 1):
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='left')
        cell.border = border

    # Add Body Rows
    csv_data = []
    current_row = header_row_idx + 1
    
    tbody = table.find('tbody')
    for tr in tbody.find_all('tr'):
        cells = tr.find_all('td')
        if not cells: continue
        
        row_data = [c.get_text(strip=True) for c in cells]
        
        if 'phase-header' in tr.get('class', []):
            # Phase Header Row
            ws.append([row_data[0]])
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers))
            cell = ws.cell(row=current_row, column=1)
            cell.fill = phase_fill
            cell.font = phase_font
            cell.border = border
            csv_data.append([f"--- {row_data[0]} ---"] + [""] * (len(headers)-1))
        else:
            # Data Row
            ws.append(row_data)
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.border = border
                # Apply conditional color for Status/Check
                if "Done" in value or "☑" in value:
                     cell.font = done_font
            csv_data.append(row_data)
        
        current_row += 1

    # Adjust column widths
    column_widths = [5, 5, 30, 50, 15, 12, 40]
    for i, width in enumerate(column_widths, 1):
        col_letter = chr(64 + i) if i <= 26 else f"A{chr(64 + i - 26)}"
        ws.column_dimensions[col_letter].width = width

    # Save files
    print(f"Saving XLSX to: {xlsx_path}")
    wb.save(xlsx_path)
    
    print(f"Saving CSV to: {csv_path}")
    df_csv = pd.DataFrame(csv_data, columns=headers)
    df_csv.to_csv(csv_path, index=False)
    
    print("Conversion complete!")

if __name__ == "__main__":
    HTML_FILE = r'c:\Users\rahul\Downloads\origin_command_execution_plan.html'
    XLSX_FILE = r'c:\Users\rahul\Downloads\origin_command_execution_plan.xlsx'
    CSV_FILE = r'c:\Users\rahul\Downloads\origin_command_execution_plan.csv'
    
    if os.path.exists(HTML_FILE):
        convert_html_to_xlsx(HTML_FILE, XLSX_FILE, CSV_FILE)
    else:
        print(f"Error: HTML file not found at {HTML_FILE}")
